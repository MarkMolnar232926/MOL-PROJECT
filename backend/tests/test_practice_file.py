"""Golden tests against the practice workbook (PLAN.md section 8).

This is the only place the hidden Answer_Key sheet is ever read.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
from openpyxl import load_workbook

from recon import InputFile, load_inputs, normalize
from recon.engine import reconcile
from recon.models import Confidence, ManualDecision, MatchStatus

# PLAN.md 6.6: the slots the Answer_Key pairs differently from FIFO.
PLAN_TIES = {
    "Swivel Chair": ({"84219511", "84247161"}, {"SN-2019-3032", "SN-2019-6040"}),
    "Work Desk": (
        {"84298003", "84214252", "84236836"},
        {"SN-2021-3808", "SN-2021-3498", "SN-2021-4894"},
    ),
    "Large Storage Cabinet": ({"84293764", "84205150"}, {"SN-2019-8043", "SN-2019-4589"}),
    "Conference Chair": ({"84264116", "84245414"}, {"SN-2025-4806", "SN-2025-3911"}),
    "Shelf Unit": ({"84287104", "84271689"}, {"SN-2024-8750", "SN-2024-5113"}),
}
# PLAN.md 8, [DECISION, phase 2]: the full attribute-based tie groups.
DETECTED_TIES = {
    "Swivel Chair": {"84219511", "84247161"},
    "Work Desk": {"84298003", "84214252", "84236836"},
    "Large Storage Cabinet": {"84293764", "84205150"},
    "Conference Chair": {"84264116", "84245414"},
    "Shelf Unit": {"84287104", "84271689", "84257423", "84229044"},
    "Coat Rack": {"84283321", "84202520"},
    "Computer Desk": {"84243853", "84235250"},
}


@pytest.fixture(scope="module")
def data():
    path = (
        Path(__file__).resolve().parents[2]
        / "tests/fixtures/Inventory_Reconciliation_Practice_1.xlsx"
    )
    if not path.exists():
        pytest.skip("practice workbook missing")
    return normalize(load_inputs([InputFile(path.name, path.read_bytes())]))


@pytest.fixture(scope="module")
def answer_key(data) -> dict[int, str]:
    """SAP excel row -> correct Asset ID (first SAP row when the key lists duplicates)."""
    path = (
        Path(__file__).resolve().parents[2]
        / "tests/fixtures/Inventory_Reconciliation_Practice_1.xlsx"
    )
    ws = load_workbook(path, read_only=True)["Answer_Key"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    i_id, i_sap = header.index("Asset ID"), header.index("SAP_Export row(s)")
    key = {}
    for row in rows[1:]:
        sap = str(row[i_sap]).strip()
        if sap != "-":
            key[int(sap.split(",")[0])] = str(row[i_id])
    return key


@pytest.fixture(scope="module")
def off(data):
    return reconcile(data, use_qr=False)


@pytest.fixture(scope="module")
def on(data):
    return reconcile(data, use_qr=True)


def rows_with(result, sheet: str, status: MatchStatus) -> set[int]:
    rows = result.physical_rows if sheet == "physical" else result.sap_rows
    return {r.excel_row for r in rows if r.match_status is status}


def test_loading_and_classification(off):
    assert off.summary.physical_total == 84
    assert off.summary.sap_total == 84
    assert all(r.sap_type for r in off.physical_rows)
    assert rows_with(off, "physical", MatchStatus.UNCLASSIFIED) == set()


def test_duplicates(off):
    assert rows_with(off, "physical", MatchStatus.DUPLICATE) == {31, 44, 83}
    assert rows_with(off, "sap", MatchStatus.DUPLICATE) == {68, 77, 79}


def test_defective_excluded(off):
    assert rows_with(off, "physical", MatchStatus.DEFECTIVE) == {42, 43, 45, 51, 57, 79, 82}


def test_pairs_and_leftovers(off):
    assert len(off.pairs) == 74
    assert rows_with(off, "physical", MatchStatus.PHYSICAL_ONLY) == set()
    assert rows_with(off, "sap", MatchStatus.SAP_ONLY) == {27, 32, 38, 50, 52, 72, 74}
    assert all(p.year_gap == 0 for p in off.pairs)


def test_location_mismatches(off):
    ids = {p.asset_id for p in off.pairs if p.location_mismatch}
    fixed = {"84208743", "84261222", "84283321", "84285415"}
    assert len(ids) == 5
    assert fixed < ids
    assert len(ids - fixed) == 1 and (ids - fixed) <= {"84287104", "84271689"}


def test_pairs_outside_plan_ties_equal_answer_key(off, answer_key):
    plan_serials = set().union(*(s for _, s in PLAN_TIES.values()))
    sap_by_row = {r.excel_row: r for r in off.sap_rows}
    outside = [p for p in off.pairs if sap_by_row[p.sap_row].serial_no not in plan_serials]
    assert len(outside) == 63
    wrong = [
        (p.sap_row, p.asset_id, answer_key[p.sap_row])
        for p in outside
        if p.asset_id != answer_key[p.sap_row]
    ]
    assert wrong == []


def test_confident_pairs_equal_answer_key(off, answer_key):
    confident = [p for p in off.pairs if p.tie_group_id is None]
    assert len(confident) == 57
    assert all(p.asset_id == answer_key[p.sap_row] for p in confident)
    assert {p.confidence for p in confident} <= {Confidence.HIGH, Confidence.MEDIUM}


def test_tie_groups(off):
    assert off.summary.tie_groups == 7
    assert off.summary.tie_slots == off.summary.tie_slots_pending == 17
    assert len(rows_with(off, "sap", MatchStatus.NEEDS_DECISION)) == 17
    by_type = {g.sap_type: g for g in off.tie_groups}
    assert set(by_type) == set(DETECTED_TIES)
    for sap_type, ids in DETECTED_TIES.items():
        g = by_type[sap_type]
        assert {c.asset_id for c in g.candidates} == ids
        assert len(g.slots) == len(ids)
        # suggestions pair only within their group, one-to-one
        suggested = [s.suggested_asset_id for s in g.slots]
        assert set(suggested) == ids
    serials = {g.sap_type: {s.serial_no for s in g.slots} for g in off.tie_groups}
    for sap_type, (ids, plan_serials) in PLAN_TIES.items():
        assert ids <= {c.asset_id for c in by_type[sap_type].candidates}
        assert plan_serials <= serials[sap_type]
    # pending slots never carry a final Asset ID
    assert all(
        r.asset_id is None for r in off.sap_rows if r.match_status is MatchStatus.NEEDS_DECISION
    )


def test_totals_reconcile(off):
    p = off.summary.physical_status_counts
    s = off.summary.sap_status_counts
    paired_p = sum(v for k, v in p.items() if k not in {"Defective – excluded", "Duplicate entry"})
    assert (paired_p, p["Defective – excluded"], p["Duplicate entry"]) == (74, 7, 3)
    paired_s = sum(v for k, v in s.items() if k not in {"SAP only", "Duplicate entry"})
    assert (paired_s, s["SAP only"], s["Duplicate entry"]) == (74, 7, 3)


def test_qr_assessment(off, on):
    for res in (off, on):  # the assessment is always computed with QR off
        qr = res.qr_assessment
        assert qr.column_present
        assert qr.parseable_pct == 100.0
        assert (qr.present_in_physical, qr.total_rows, qr.present_pct) == (77, 84, 91.7)
        assert qr.agreement_outside_ties_pct == 100.0
        assert qr.looks_reliable


def test_qr_on_matches_answer_key(on, answer_key):
    assert len(on.pairs) == 74
    assert all(p.asset_id == answer_key[p.sap_row] for p in on.pairs)
    assert all(p.confidence is Confidence.HIGH_QR for p in on.pairs)
    assert on.summary.tie_slots_pending == 0
    assert rows_with(on, "sap", MatchStatus.SAP_ONLY) == {27, 32, 38, 50, 52, 72, 74}
    assert all(
        "QR references missing asset" in " ".join(r.notes)
        for r in on.sap_rows
        if r.match_status is MatchStatus.SAP_ONLY
    )
    assert on.summary.location_mismatches == 5


def test_qr_off_validation_column(off):
    confident = {p.sap_row for p in off.pairs if p.tie_group_id is None}
    assert all(r.qr_agrees for r in off.sap_rows if r.excel_row in confident)


def test_resolving_ties_with_answer_key_gives_full_match(data, off, answer_key):
    now = dt.datetime(2026, 1, 1)
    decisions = [
        ManualDecision(
            group_id=g.group_id,
            sap_row=s.sap_row,
            asset_id=answer_key[s.sap_row],
            proposed_asset_id=s.suggested_asset_id,
            decided_at=now,
        )
        for g in off.tie_groups
        for s in g.slots
    ]
    res = reconcile(data, decisions=decisions)
    assert res.summary.tie_slots_pending == 0
    assert res.summary.manual_decisions == 17
    assert all(p.asset_id == answer_key[p.sap_row] for p in res.pairs)
    sap_by_row = {r.excel_row: r for r in res.sap_rows}
    for d in decisions:
        assert sap_by_row[d.sap_row].asset_id == d.asset_id
        assert sap_by_row[d.sap_row].match_status is MatchStatus.MANUAL


def test_deactivated_warnings(off):
    flagged = {d.asset_id for d in off.discrepancies if d.kind.value == "Deactivated warning"}
    assert flagged == {"84213492", "84200684", "84241548"}
