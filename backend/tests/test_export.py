import datetime as dt
import io
from collections import Counter

import pytest
from conftest import PRACTICE_FILE, make_input, phys, sap
from openpyxl import load_workbook

from recon import InputFile, load_inputs, normalize
from recon.engine import accept_suggestions, reconcile
from recon.export import SAP_EXTRA, export_filename, export_workbook
from recon.models import MatchStatus

SHEETS = [
    "SAP_Export_Completed",
    "Physical_Inventory_Annotated",
    "Discrepancies",
    "Summary",
    "Decisions_Log",
]


@pytest.fixture(scope="module")
def practice():
    if not PRACTICE_FILE.exists():
        pytest.skip("practice workbook missing")
    data = normalize(load_inputs([InputFile(PRACTICE_FILE.name, PRACTICE_FILE.read_bytes())]))
    pending = reconcile(data)
    accepted = reconcile(data, decisions=accept_suggestions(pending))
    return data, pending, accepted


def roundtrip(data, result):
    return load_workbook(io.BytesIO(export_workbook(data.loaded, result)))


def table(ws):
    rows = list(ws.iter_rows(values_only=True))
    return list(rows[0]), rows[1:]


def summary_values(wb):
    return {r[1]: r[2] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}


def test_filename():
    assert export_filename(dt.datetime(2026, 1, 2, 3, 4, 5)) == "reconciled_20260102_030405.xlsx"


def test_sheets_and_layout(practice):
    data, pending, _ = practice
    wb = roundtrip(data, pending)
    assert wb.sheetnames == SHEETS
    for ws in wb:
        assert ws.freeze_panes == "A2"
        assert ws["A1"].font.name == "Arial" and ws["A1"].font.b
        for row in ws.iter_rows():
            for cell in row:
                assert not (isinstance(cell.value, str) and cell.value.startswith("="))


def test_sap_sheet_keeps_original_columns_and_values(practice):
    data, pending, _ = practice
    headers, rows = table(roundtrip(data, pending)["SAP_Export_Completed"])
    source = load_workbook(PRACTICE_FILE, read_only=True)["SAP_Export"]
    src_rows = list(source.iter_rows(values_only=True))
    n = len(src_rows[0])
    assert headers == list(src_rows[0]) + SAP_EXTRA
    assert len(rows) == 84
    for out, src in zip(rows, src_rows[1:], strict=True):
        assert out[1:n] == tuple(src[1:n])  # everything except Asset ID untouched


def test_asset_ids_in_the_right_rows(practice):
    data, pending, accepted = practice
    for result, empty_pending in ((pending, True), (accepted, False)):
        ws = roundtrip(data, result)["SAP_Export_Completed"]
        headers, rows = table(ws)
        ai, si = headers.index("Asset ID"), headers.index("Match status")
        by_row = {r.excel_row: r for r in result.sap_rows}
        for excel_row, row in enumerate(rows, start=2):
            expected = by_row[excel_row].asset_id
            assert row[ai] == (int(expected) if expected else None)
            if row[si] == MatchStatus.NEEDS_DECISION.value:
                assert empty_pending and row[ai] is None  # never the provisional suggestion
        assert ws.cell(2, ai + 1).fill.fgColor.rgb == "FFFFFFCC"
    _, rows = table(roundtrip(data, pending)["SAP_Export_Completed"])
    assert sum(r[0] is not None for r in rows) == 57
    _, rows = table(roundtrip(data, accepted)["SAP_Export_Completed"])
    assert sum(r[0] is not None for r in rows) == 74


def test_physical_sheet(practice):
    data, pending, _ = practice
    headers, rows = table(roundtrip(data, pending)["Physical_Inventory_Annotated"])
    assert headers[-4:] == ["Match status", "Matched SAP row", "SAP Item Name", "Notes"]
    assert len(rows) == 84
    counts = Counter(r[headers.index("Match status")] for r in rows)
    assert counts == Counter(pending.summary.physical_status_counts)


def test_discrepancies_and_summary(practice):
    data, pending, accepted = practice
    wb = roundtrip(data, pending)
    headers, rows = table(wb["Discrepancies"])
    assert headers[0] == "Issue"
    assert Counter(r[0] for r in rows) == Counter(d.kind.value for d in pending.discrepancies)
    assert all(r[3] is None for r in rows if r[0] == "Unresolved tie")
    s = summary_values(wb)
    assert s["Tie slots still needing a decision"] == 17
    assert s["Rule-set version"] == "1"
    assert summary_values(roundtrip(data, accepted))["Manual decisions"] == 17


def test_decisions_log(practice):
    data, pending, accepted = practice
    _, rows = table(roundtrip(data, pending)["Decisions_Log"])
    assert rows == []
    headers, rows = table(roundtrip(data, accepted)["Decisions_Log"])
    assert headers[-1] == "Decided at (UTC)"
    assert len(rows) == 17
    assert all(r[3] == r[4] for r in rows)  # accepted suggestions
    assert all(isinstance(r[5], dt.datetime) for r in rows)


def test_export_without_qr_column_and_with_none_decision():
    data = make_input(
        [phys(1, "bin", "bin, grey"), phys(2, "bin", "bin, grey")],
        [sap("Trash Bin", "Grey", 30, "SN-2020-1"), sap("Trash Bin", "Grey", 30, "SN-2020-2")],
        drop_sap=("QR Code",),
    )
    result = reconcile(data)
    (g,) = result.tie_groups
    decisions = accept_suggestions(result)
    decisions[0] = decisions[0].model_copy(update={"asset_id": None})
    result = reconcile(data, decisions=decisions)
    wb = roundtrip(data, result)
    headers, rows = table(wb["SAP_Export_Completed"])
    assert "QR Code" not in headers and headers[-len(SAP_EXTRA) :] == SAP_EXTRA
    assert rows[0][0] is None and rows[0][headers.index("Match status")] == "SAP only"
    _, log = table(wb["Decisions_Log"])
    assert log[0][3] == "none (left unmatched)"
