import datetime as dt
import io

import pytest
from conftest import PHYSICAL_HEADERS, PRACTICE_FILE, SAP_HEADERS, build_workbook, phys, sap
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import api.main as api_main
from api.main import create_app
from api.sessions import InMemorySessionStore


@pytest.fixture
def client():
    return TestClient(create_app(store=InMemorySessionStore()))


@pytest.fixture(scope="module")
def practice_bytes():
    if not PRACTICE_FILE.exists():
        pytest.skip("practice workbook missing")
    return PRACTICE_FILE.read_bytes()


@pytest.fixture
def session(client, practice_bytes):
    r = client.post("/api/sessions", files={"workbook": ("practice.xlsx", practice_bytes)})
    assert r.status_code == 201, r.text
    return r.json()


def upload_two(client, physical_rows, sap_rows):
    p = build_workbook({"Stocktake": [PHYSICAL_HEADERS, *physical_rows]})
    s = build_workbook({"Export": [SAP_HEADERS, *sap_rows]})
    return client.post("/api/sessions", files={"physical": ("p.xlsx", p), "sap": ("s.xlsx", s)})


def two_chairs():
    return (
        [
            phys(11111111, "chair", "on wheels, black", date=(2019, 5, 1)),
            phys(22222222, "chair", "on wheels, black", date=(2019, 2, 1)),
        ],
        [
            sap("Swivel Chair", "Black", 60, "SN-2019-9000", qr="INV0011111111"),
            sap("Swivel Chair", "Black", 60, "SN-2019-1000", qr="INV0022222222"),
        ],
    )


def export_rows(client, sid):
    r = client.get(f"/api/sessions/{sid}/export")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["SAP_Export_Completed"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    return {i: dict(zip(header, row, strict=True)) for i, row in enumerate(rows[1:], start=2)}


# --- basics ---------------------------------------------------------------------------------


def test_health_and_config(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    body = client.get("/api/config").json()
    assert len(body["type_rules"]) == 25
    assert {loc["building"] for loc in body["locations"]} == {"RVS", "LKS"}
    assert body["cost_weights"]["year_gap"] == 1000


def test_openapi_lists_endpoints(client):
    paths = client.get("/openapi.json").json()["paths"]
    assert set(paths) == {
        "/api/health",
        "/api/config",
        "/api/sessions",
        "/api/sessions/{session_id}",
        "/api/sessions/{session_id}/result",
        "/api/sessions/{session_id}/ties/{group_id}",
        "/api/sessions/{session_id}/export",
    }


# --- upload ---------------------------------------------------------------------------------


def test_upload_one_workbook(session):
    assert [
        (d["source"], d["sheet_name"], d["detected_by"], d["row_count"])
        for d in session["detected"]
    ] == [
        ("physical", "Physical_Inventory", "sheet_name", 84),
        ("sap", "SAP_Export", "sheet_name", 84),
    ]
    assert session["summary"]["pairs"] == 74
    assert session["summary"]["tie_slots_pending"] == 17


def test_upload_two_files(client):
    r = upload_two(client, *two_chairs())
    assert r.status_code == 201, r.text
    detected = {d["source"]: d for d in r.json()["detected"]}
    assert detected["physical"]["file_name"] == "p.xlsx"
    assert detected["sap"]["detected_by"] == "header_signature"


@pytest.mark.parametrize(
    "files",
    [
        {"workbook": ("notes.csv", b"a,b\n1,2")},
        {"workbook": ("fake.xlsx", b"not a zip")},
        {"physical": ("p.xlsx", build_workbook({"x": [["a"]]}))},  # sap missing
        {},
    ],
)
def test_bad_uploads_are_422(client, files):
    r = client.post("/api/sessions", files=files) if files else client.post("/api/sessions")
    assert r.status_code == 422
    err = r.json()["error"]
    assert err["code"] in {"invalid_file", "invalid_request"}
    assert "Traceback" not in r.text


def test_both_modes_at_once_is_rejected(client, practice_bytes):
    files = {
        "workbook": ("a.xlsx", practice_bytes),
        "physical": ("b.xlsx", practice_bytes),
        "sap": ("c.xlsx", practice_bytes),
    }
    r = client.post("/api/sessions", files=files)
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_file"


def test_missing_columns_are_listed(client):
    drop = SAP_HEADERS.index("Serial No.")
    sap_sheet = [[v for i, v in enumerate(r) if i != drop] for r in [SAP_HEADERS]]
    wb = build_workbook({"Physical_Inventory": [PHYSICAL_HEADERS], "SAP_Export": sap_sheet})
    r = client.post("/api/sessions", files={"workbook": ("b.xlsx", wb)})
    assert r.status_code == 422
    err = r.json()["error"]
    assert err["code"] == "missing_columns"
    assert err["details"][0]["missing_columns"] == ["Serial No."]


def test_upload_limit(client, monkeypatch):
    monkeypatch.setattr(api_main, "MAX_UPLOAD_BYTES", 10)
    r = client.post("/api/sessions", files={"workbook": ("big.xlsx", b"x" * 11)})
    assert r.status_code == 413 and r.json()["error"]["code"] == "file_too_large"


# --- sessions -------------------------------------------------------------------------------


def test_unknown_session_and_group(client, session):
    assert client.get("/api/sessions/nope/result").status_code == 404
    r = client.put(f"/api/sessions/{session['session_id']}/ties/nope", json={"assignments": []})
    assert r.status_code == 404 and r.json()["error"]["code"] == "tie_group_not_found"


def test_result(client, session):
    res = client.get(f"/api/sessions/{session['session_id']}/result").json()
    assert res["session_id"] == session["session_id"]
    assert len(res["sap_rows"]) == len(res["physical_rows"]) == 84
    assert len(res["tie_groups"]) == 7
    assert "qr_assessment" not in res
    assert all(r["qr_code"] for r in res["sap_rows"])  # kept as a label
    assert all(g["pending_slots"] == len(g["slots"]) for g in res["tie_groups"])


def test_delete_session(client, session):
    sid = session["session_id"]
    assert client.delete(f"/api/sessions/{sid}").status_code == 204
    assert client.get(f"/api/sessions/{sid}/result").status_code == 404


def test_sessions_expire_after_idle_time():
    now = [dt.datetime(2026, 1, 1, tzinfo=dt.UTC)]
    store = InMemorySessionStore(idle_ttl=dt.timedelta(hours=2), clock=lambda: now[0])
    client = TestClient(create_app(store=store))
    sid = upload_two(client, *two_chairs()).json()["session_id"]
    now[0] += dt.timedelta(hours=1, minutes=59)
    assert client.get(f"/api/sessions/{sid}/result").status_code == 200  # access renews
    now[0] += dt.timedelta(hours=1, minutes=59)
    assert client.get(f"/api/sessions/{sid}/result").status_code == 200
    now[0] += dt.timedelta(hours=2, minutes=1)
    assert client.get(f"/api/sessions/{sid}/result").status_code == 404
    assert len(store) == 0


# --- ties -----------------------------------------------------------------------------------


def _group(client, sid):
    (group,) = client.get(f"/api/sessions/{sid}/result").json()["tie_groups"]
    return group


def test_tie_decision_flow(client):
    sid = upload_two(client, *two_chairs()).json()["session_id"]
    g = _group(client, sid)
    gid = g["group_id"]
    assert {s["sap_row"]: s["suggested_asset_id"] for s in g["slots"]} == {
        2: "11111111",
        3: "22222222",
    }
    # pick the opposite of the suggestion for one slot
    r = client.put(
        f"/api/sessions/{sid}/ties/{gid}",
        json={"assignments": [{"sap_row": 2, "physical_asset_id": "22222222"}]},
    )
    assert r.status_code == 200, r.text
    slots = {s["sap_row"]: s for s in r.json()["slots"]}
    assert slots[2]["decided"] and slots[2]["chosen_asset_id"] == "22222222"
    assert slots[3]["suggested_asset_id"] == "11111111"  # remaining slot re-suggested
    assert r.json()["pending_slots"] == 1
    # one-to-one: the other slot cannot take the same unit
    r = client.put(
        f"/api/sessions/{sid}/ties/{gid}",
        json={"assignments": [{"sap_row": 3, "physical_asset_id": "22222222"}]},
    )
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_tie_decision"
    # not a candidate
    r = client.put(
        f"/api/sessions/{sid}/ties/{gid}",
        json={"assignments": [{"sap_row": 3, "physical_asset_id": "99999999"}]},
    )
    assert r.status_code == 422
    # "none" is allowed
    r = client.put(
        f"/api/sessions/{sid}/ties/{gid}",
        json={"assignments": [{"sap_row": 3, "physical_asset_id": None}]},
    )
    assert r.status_code == 200 and r.json()["pending_slots"] == 0
    res = client.get(f"/api/sessions/{sid}/result").json()
    assert {r["excel_row"]: r["match_status"] for r in res["sap_rows"]} == {
        2: "Manually resolved",
        3: "SAP only",
    }
    assert len(res["decisions"]) == 2
    # reset
    r = client.delete(f"/api/sessions/{sid}/ties/{gid}")
    assert r.status_code == 200 and r.json()["pending_slots"] == 2
    assert client.get(f"/api/sessions/{sid}/result").json()["decisions"] == []


# --- export -------------------------------------------------------------------------------


def test_export_roundtrip(client, session):
    sid = session["session_id"]
    r = client.get(f"/api/sessions/{sid}/export")
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="reconciled_' in r.headers["content-disposition"]
    res = client.get(f"/api/sessions/{sid}/result").json()
    rows = export_rows(client, sid)
    for sap_row in res["sap_rows"]:
        out = rows[sap_row["excel_row"]]
        expected = sap_row["asset_id"]
        assert out["Asset ID"] == (int(expected) if expected else None)
        assert out["Match status"] == sap_row["match_status"]
    pending = [r for r in rows.values() if r["Match status"] == "Needs decision"]
    assert len(pending) == 17 and all(r["Asset ID"] is None for r in pending)

    # resolve one group by accepting its suggestion and export again
    g = res["tie_groups"][0]
    body = {
        "assignments": [
            {"sap_row": s["sap_row"], "physical_asset_id": s["suggested_asset_id"]}
            for s in g["slots"]
        ]
    }
    assert client.put(f"/api/sessions/{sid}/ties/{g['group_id']}", json=body).status_code == 200
    rows = export_rows(client, sid)
    for s in g["slots"]:
        assert rows[s["sap_row"]]["Asset ID"] == int(s["suggested_asset_id"])
        assert rows[s["sap_row"]]["Match status"] == "Manually resolved"
