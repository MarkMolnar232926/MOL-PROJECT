"""The export: the uploaded incoming file with only its Asset ID cells filled in (section 6)."""

import hashlib
import io

import pytest
from conftest import INCOMING_FILE, ORIGINAL_FILE, PHYSICAL_HEADERS, SAP_HEADERS, phys, sap
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from api.main import create_app
from api.sessions import InMemorySessionStore
from recon.writeback import export_filename


@pytest.fixture
def client():
    return TestClient(create_app(store=InMemorySessionStore()))


def start(client, original: bytes, incoming: bytes, incoming_name="sap_export.xlsx") -> str:
    sid = client.post("/api/sessions").json()["session_id"]
    r = client.post(f"/api/sessions/{sid}/original", files={"file": ("o.xlsx", original)})
    assert r.status_code == 200, r.text
    r = client.post(f"/api/sessions/{sid}/incoming", files={"file": (incoming_name, incoming)})
    assert r.status_code == 200, r.text
    assert client.post(f"/api/sessions/{sid}/match").status_code == 200
    return sid


def no_pair(client, sid, row, reason="missing_asset"):
    r = client.put(
        f"/api/sessions/{sid}/incoming/{row}/assignment",
        json={"asset_id": None, "reason": reason},
    )
    assert r.status_code == 200, r.text


def all_cells(data: bytes) -> dict[str, dict[tuple[int, int], object]]:
    wb = load_workbook(io.BytesIO(data))
    return {
        ws.title: {
            (c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None
        }
        for ws in wb.worksheets
    }


@pytest.fixture(scope="module")
def samples():
    if not (ORIGINAL_FILE.exists() and INCOMING_FILE.exists()):
        pytest.skip("sample files missing from tests/fixtures")
    return ORIGINAL_FILE.read_bytes(), INCOMING_FILE.read_bytes()


def test_filename():
    assert export_filename("sap_export.xlsx") == "sap_export_asset_id.xlsx"
    assert export_filename("Bútorok 2026.XLSX") == "Bútorok 2026_asset_id.xlsx"
    assert export_filename("macro.xlsm") == "macro_asset_id.xlsm"


def test_export_is_409_while_items_are_unresolved(client, samples):
    sid = start(client, *samples)
    r = client.get(f"/api/sessions/{sid}/export")
    assert r.status_code == 409
    err = r.json()["error"]
    assert err["code"] == "unresolved_items"
    assert err["details"][0]["rows"] == [8, 14, 16, 20]
    for row in (8, 14, 16):
        no_pair(client, sid, row)
    assert client.get(f"/api/sessions/{sid}/export").status_code == 409
    no_pair(client, sid, 20, "new_asset")
    assert client.get(f"/api/sessions/{sid}/export").status_code == 200


def test_export_integrity_on_the_sample(client, samples):
    original_hash = hashlib.sha256(ORIGINAL_FILE.read_bytes()).hexdigest()
    incoming_hash = hashlib.sha256(INCOMING_FILE.read_bytes()).hexdigest()
    original, incoming = samples
    sid = start(client, original, incoming)
    # rows 8 and 16: take the best (location differs) unit by hand; 14 and 20: no pair
    for row, asset_id in ((8, "84277502"), (16, "84285415")):
        r = client.put(
            f"/api/sessions/{sid}/incoming/{row}/assignment", json={"asset_id": asset_id}
        )
        assert r.status_code == 200, r.text
    no_pair(client, sid, 14)
    no_pair(client, sid, 20)
    result = client.get(f"/api/sessions/{sid}/result").json()

    r = client.get(f"/api/sessions/{sid}/export")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="sap_export_asset_id.xlsx"' in r.headers["content-disposition"]

    before, after = all_cells(incoming), all_cells(r.content)
    assert list(after) == list(before) == ["Munka1"]
    col = SAP_HEADERS.index("Asset ID") + 1
    # every cell outside the Asset ID data cells is unchanged
    strip = lambda cells: {k: v for k, v in cells.items() if k[1] != col or k[0] == 1}  # noqa: E731
    assert strip(after["Munka1"]) == strip(before["Munka1"])
    # the Asset IDs are the matching result: 18 filled, the two "no pair" rows empty
    written = {k[0]: v for k, v in after["Munka1"].items() if k[1] == col and k[0] > 1}
    expected = {x["item"]["excel_row"]: int(x["asset_id"]) for x in result["rows"] if x["asset_id"]}
    assert written == expected
    assert len(written) == 18 and 14 not in written and 20 not in written
    # the formatting of the Asset ID column survives
    ws_in = load_workbook(io.BytesIO(incoming))["Munka1"]
    ws_out = load_workbook(io.BytesIO(r.content))["Munka1"]
    for row in range(1, ws_in.max_row + 1):
        a, b = ws_in.cell(row, col), ws_out.cell(row, col)
        assert (a.fill.fgColor.rgb, a.font.b, a.number_format) == (
            b.fill.fgColor.rgb,
            b.font.b,
            b.number_format,
        )
    # neither uploaded file changed on disk
    assert hashlib.sha256(ORIGINAL_FILE.read_bytes()).hexdigest() == original_hash
    assert hashlib.sha256(INCOMING_FILE.read_bytes()).hexdigest() == incoming_hash


def rich_incoming() -> bytes:
    """An incoming workbook with a notes sheet, a formula, styles and odd column widths."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(SAP_HEADERS)
    ws.append(sap("Swivel Chair", "Black", 60, "SN-1"))
    ws.append(sap("Swivel Chair", "Black", 60, "SN-2", remarks="old label"))
    ws.append(sap("Trash Bin", "Grey", 30, "SN-3"))
    ws["A3"] = 12345678  # pre-filled Asset ID on a row that will be matched
    ws["A4"] = 87654321  # pre-filled Asset ID on a row that gets "no pair"
    ws["N1"] = "Check"
    ws["N2"] = "=G2*2"
    for c in ws[1]:
        c.font = Font(bold=True)
    ws["A2"].fill = PatternFill("solid", fgColor="FFFFFFCC")
    ws.column_dimensions["B"].width = 42
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Do not touch"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_writes_only_asset_ids(client):
    original = Workbook()
    ws = original.active
    ws.append(PHYSICAL_HEADERS)
    ws.append(phys(11111111, "chair", "on wheels, black", date=(2019, 1, 1)))
    ws.append(phys(22222222, "chair", "on wheels, black", date=(2020, 1, 1)))
    buf = io.BytesIO()
    original.save(buf)
    incoming = rich_incoming()
    sid = start(client, buf.getvalue(), incoming, "Bútor lista.xlsx")
    no_pair(client, sid, 4, "new_asset")  # the trash bin

    r = client.get(f"/api/sessions/{sid}/export")
    assert r.status_code == 200, r.text
    assert "filename*=UTF-8''B%C3%BAtor%20lista_asset_id.xlsx" in r.headers["content-disposition"]
    out = load_workbook(io.BytesIO(r.content))
    assert out.sheetnames == ["Export", "Notes"]
    ws = out["Export"]
    assert [ws.cell(r, 1).value for r in (2, 3, 4)] == [22222222, 11111111, None]
    assert ws["N2"].value == "=G2*2"  # formulas stay formulas
    assert ws["A1"].font.b and ws["A2"].fill.fgColor.rgb == "FFFFFFCC"
    assert ws.column_dimensions["B"].width == 42
    assert ws.freeze_panes == "A2"
    assert out["Notes"]["A1"].value == "Do not touch"
    before, after = all_cells(incoming), all_cells(r.content)
    for sheet in before:
        a = {k: v for k, v in before[sheet].items() if not (sheet == "Export" and k[1] == 1)}
        b = {k: v for k, v in after[sheet].items() if not (sheet == "Export" and k[1] == 1)}
        assert a == b


def test_asset_id_column_found_by_header_not_position(client):
    headers = [*SAP_HEADERS[1:], SAP_HEADERS[0]]  # Asset ID last
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    row = sap("Swivel Chair", "Black", 60, "SN-1")
    ws.append([*row[1:], row[0]])
    buf = io.BytesIO()
    wb.save(buf)
    original = Workbook()
    original.active.append(PHYSICAL_HEADERS)
    original.active.append(phys(11111111, "chair", "on wheels, black"))
    obuf = io.BytesIO()
    original.save(obuf)
    sid = start(client, obuf.getvalue(), buf.getvalue())
    r = client.get(f"/api/sessions/{sid}/export")
    out = load_workbook(io.BytesIO(r.content)).active
    assert out.cell(2, len(headers)).value == 11111111
    assert out.cell(2, 1).value == "> Movable Furniture"
