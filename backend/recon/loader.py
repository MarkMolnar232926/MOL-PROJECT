"""Read the uploaded workbook(s) and locate the Physical_Inventory and SAP_Export sheets.

Input is either one workbook holding both sheets or two files (one per source).
Each source is found by sheet name first, then by header signature. Sheets listed in
``ignored_sheets`` (the hidden Answer_Key) are never opened.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from typing import Literal

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .config import AppConfig, Source, SourceSpec, default_config, normalize_header
from .errors import InvalidFileError, MissingColumnsError, SheetDetectionError

ALLOWED_EXTENSIONS = (".xlsx", ".xlsm")
SOURCES: tuple[Source, ...] = ("physical", "sap")
SOURCE_LABELS = {"physical": "Physical_Inventory", "sap": "SAP_Export"}
EXCEL_ROW = "excel_row"

DetectedBy = Literal["sheet_name", "header_signature"]


@dataclass
class InputFile:
    name: str
    data: bytes
    role: Source | None = None  # the upload slot the file came from, if any


@dataclass
class SheetData:
    source: Source
    file_name: str
    sheet_name: str
    detected_by: DetectedBy
    headers: list[str]  # original header texts, in sheet order
    column_map: dict[str, str]  # canonical key -> original header
    frame: pd.DataFrame  # one column per original header + EXCEL_ROW
    missing_optional: list[str] = field(default_factory=list)
    column_fills: dict[str, str] = field(default_factory=dict)  # header -> ARGB of whole column

    @property
    def row_count(self) -> int:
        return len(self.frame)


@dataclass
class LoadedInput:
    physical: SheetData
    sap: SheetData
    warnings: list[str] = field(default_factory=list)

    def sheet(self, source: Source) -> SheetData:
        return self.physical if source == "physical" else self.sap


@dataclass
class _Candidate:
    file: InputFile
    workbook_index: int
    sheet_name: str
    headers: list[str]


def _open_workbook(file: InputFile):
    if not file.name.lower().endswith(ALLOWED_EXTENSIONS):
        raise InvalidFileError(
            f"'{file.name}' is not an Excel workbook. Upload a .xlsx or .xlsm file.",
            [{"file": file.name, "problem": "unsupported file type"}],
        )
    if not zipfile.is_zipfile(io.BytesIO(file.data)):
        raise InvalidFileError(
            f"'{file.name}' could not be read as an Excel workbook (the file looks damaged "
            "or is an old .xls file saved with a new extension).",
            [{"file": file.name, "problem": "not a valid xlsx archive"}],
        )
    try:
        return load_workbook(io.BytesIO(file.data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        raise InvalidFileError(
            f"'{file.name}' could not be opened: {exc}",
            [{"file": file.name, "problem": "unreadable workbook"}],
        ) from None


def _header_row(ws: Worksheet) -> list[str]:
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = ["" if v is None else str(v).strip() for v in row]
        while headers and headers[-1] == "":
            headers.pop()
        return headers
    return []


def _map_columns(headers: list[str], spec: SourceSpec) -> dict[str, str]:
    by_norm: dict[str, str] = {}
    for h in headers:
        by_norm.setdefault(normalize_header(h), h)
    mapping: dict[str, str] = {}
    for key, col in spec.columns.items():
        for alias in col.headers:
            if (hit := by_norm.get(normalize_header(alias))) is not None:
                mapping[key] = hit
                break
    return mapping


def _missing_required(mapping: dict[str, str], spec: SourceSpec) -> list[str]:
    return [spec.columns[k].display_name for k in spec.required_keys if k not in mapping]


def _fill_rgb(cell) -> str | None:
    fill = getattr(cell, "fill", None)
    if fill is None or not fill.fill_type or fill.fgColor is None:
        return None
    rgb = fill.fgColor.rgb
    return rgb if isinstance(rgb, str) else None


def _read_frame(ws: Worksheet, headers: list[str]) -> tuple[pd.DataFrame, dict[str, str]]:
    """Read data rows; also return fills shared by every data cell of a column."""
    labels = [h or f"Column {i + 1}" for i, h in enumerate(headers)]
    seen: dict[str, int] = {}
    for i, label in enumerate(labels):  # make duplicate header texts unique
        if label in seen:
            seen[label] += 1
            labels[i] = f"{label} ({seen[label]})"
        else:
            seen[label] = 1
    width = len(labels)
    records: list[list] = []
    rows: list[int] = []
    fills: list[set[str | None]] = [set() for _ in range(width)]
    for excel_row, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        cells = list(cells[:width])
        values = [c.value for c in cells] + [None] * max(0, width - len(cells))
        values = [v.strip() if isinstance(v, str) else v for v in values]
        values = [None if v == "" else v for v in values]
        if all(v is None for v in values):
            continue
        records.append(values)
        rows.append(excel_row)
        for i in range(width):
            fills[i].add(_fill_rgb(cells[i]) if i < len(cells) else None)
    frame = pd.DataFrame(records, columns=labels, dtype=object)
    frame[EXCEL_ROW] = rows
    column_fills = {
        label: next(iter(f)) for label, f in zip(labels, fills, strict=True) if len(f) == 1
    }
    return frame, {k: v for k, v in column_fills.items() if v}


def _choose(
    source: Source, candidates: list[_Candidate], cfg: AppConfig, taken: set[tuple[int, str]]
) -> tuple[_Candidate, DetectedBy] | dict:
    """Pick the best sheet for ``source``; on failure return an error-detail dict."""
    spec = cfg.columns.source(source)
    wanted_names = {normalize_header(n) for n in spec.sheet_names}

    def role_rank(c: _Candidate) -> int:
        return 0 if c.file.role == source else (1 if c.file.role is None else 2)

    pool = sorted(
        (c for c in candidates if (c.workbook_index, c.sheet_name) not in taken), key=role_rank
    )
    by_name = [c for c in pool if normalize_header(c.sheet_name) in wanted_names]
    # A sheet with the expected name is authoritative, even if columns are missing.
    for c in by_name:
        missing = _missing_required(_map_columns(c.headers, spec), spec)
        if not missing:
            return c, "sheet_name"
    if by_name:
        c = by_name[0]
        return {
            "source": source,
            "file": c.file.name,
            "sheet": c.sheet_name,
            "missing_columns": _missing_required(_map_columns(c.headers, spec), spec),
        }
    for c in pool:
        if not _missing_required(_map_columns(c.headers, spec), spec):
            return c, "header_signature"
    # Nothing complete: report the closest partial match, if any.
    best, best_found = None, 0
    for c in pool:
        missing = _missing_required(_map_columns(c.headers, spec), spec)
        found = len(spec.required_keys) - len(missing)
        if found > best_found:
            best, best_found = c, found
    if best is not None and best_found * 2 >= len(spec.required_keys):
        return {
            "source": source,
            "file": best.file.name,
            "sheet": best.sheet_name,
            "missing_columns": _missing_required(_map_columns(best.headers, spec), spec),
        }
    return {"source": source, "file": None, "sheet": None, "missing_columns": None}


def load_inputs(files: list[InputFile], cfg: AppConfig | None = None) -> LoadedInput:
    """Load one workbook with both sheets, or two files with one source each."""
    cfg = cfg or default_config()
    if not 1 <= len(files) <= 2:
        raise InvalidFileError("Upload either one workbook or two files (physical + SAP).")
    ignored = {normalize_header(n) for n in cfg.columns.ignored_sheets}

    workbooks = [_open_workbook(f) for f in files]
    try:
        candidates: list[_Candidate] = []
        for idx, (file, wb) in enumerate(zip(files, workbooks, strict=True)):
            for name in wb.sheetnames:
                if normalize_header(name) in ignored:
                    continue
                candidates.append(_Candidate(file, idx, name, _header_row(wb[name])))

        chosen: dict[Source, tuple[_Candidate, DetectedBy]] = {}
        problems: list[dict] = []
        taken: set[tuple[int, str]] = set()
        for source in SOURCES:
            result = _choose(source, candidates, cfg, taken)
            if isinstance(result, dict):
                problems.append(result)
            else:
                chosen[source] = result
                taken.add((result[0].workbook_index, result[0].sheet_name))

        if problems:
            _raise_detection_problems(problems, candidates)

        warnings: list[str] = []
        sheets: dict[Source, SheetData] = {}
        for source, (cand, detected_by) in chosen.items():
            spec = cfg.columns.source(source)
            if cand.file.role not in (None, source):
                warnings.append(
                    f"{SOURCE_LABELS[source]} was found in '{cand.file.name}', which was uploaded "
                    f"as the {cand.file.role} file. The files may have been swapped."
                )
            ws = workbooks[cand.workbook_index][cand.sheet_name]
            frame, column_fills = _read_frame(ws, cand.headers)
            labels = [c for c in frame.columns if c != EXCEL_ROW]
            mapping = _map_columns(labels, spec)
            sheets[source] = SheetData(
                source=source,
                file_name=cand.file.name,
                sheet_name=cand.sheet_name,
                detected_by=detected_by,
                headers=labels,
                column_map=mapping,
                frame=frame,
                column_fills=column_fills,
                missing_optional=[
                    c.display_name for k, c in spec.columns.items() if k not in mapping
                ],
            )
        return LoadedInput(physical=sheets["physical"], sap=sheets["sap"], warnings=warnings)
    finally:
        for wb in workbooks:
            wb.close()


def _raise_detection_problems(problems: list[dict], candidates: list[_Candidate]) -> None:
    seen = [f"{c.file.name} › {c.sheet_name}" for c in candidates]
    lines: list[str] = []
    details: list[dict] = []
    missing_any = False
    for p in problems:
        label = SOURCE_LABELS[p["source"]]
        if p["missing_columns"]:
            missing_any = True
            lines.append(
                f"{label} (sheet '{p['sheet']}' in '{p['file']}') is missing required "
                f"column(s): {', '.join(p['missing_columns'])}."
            )
        else:
            lines.append(f"No {label} sheet found (looked by sheet name and by column headers).")
        details.append({**p, "sheets_seen": seen})
    error_cls = MissingColumnsError if missing_any else SheetDetectionError
    raise error_cls(" ".join(lines), details)
