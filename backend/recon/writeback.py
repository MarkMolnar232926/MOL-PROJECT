"""The one download: the uploaded incoming workbook with its Asset ID cells filled in.

The uploaded file is opened with openpyxl and only the Asset ID cells of the matched sheet's
data rows are written; every other cell, sheet, style and column is left as it was. Rows with
no pair get an empty Asset ID. The original inventory is never part of an export.
"""

from __future__ import annotations

import io
from pathlib import PurePath

from openpyxl import load_workbook

from .errors import ReconError
from .intake import SourceUpload
from .match_models import MatchResult


class UnresolvedItemsError(ReconError):
    """Export is only possible once every incoming row is resolved."""

    code = "unresolved_items"


def export_filename(uploaded_name: str) -> str:
    """``<original name>_asset_id.xlsx`` (``.xlsm`` stays ``.xlsm`` so macros remain valid)."""
    p = PurePath(uploaded_name or "incoming.xlsx")
    ext = ".xlsm" if p.suffix.lower() == ".xlsm" else ".xlsx"
    return f"{p.stem}_asset_id{ext}"


def _cell_value(asset_id: str | None):
    if asset_id is None:
        return None
    return int(asset_id) if asset_id.isdigit() else asset_id


def asset_id_column(upload: SourceUpload) -> int:
    """1-based column of the Asset ID header in the incoming sheet."""
    sheet = upload.sheet
    return sheet.headers.index(sheet.column_map["asset_id"]) + 1


def write_asset_ids(upload: SourceUpload, result: MatchResult) -> bytes:
    """The uploaded incoming workbook with the Asset IDs of ``result`` written in."""
    unresolved = [r.item.excel_row for r in result.rows if not r.resolved]
    if unresolved:
        raise UnresolvedItemsError(
            f"{len(unresolved)} item(s) still need a decision before the export.",
            [{"unresolved": len(unresolved), "rows": unresolved}],
        )
    keep_vba = upload.file.name.lower().endswith(".xlsm")
    wb = load_workbook(io.BytesIO(upload.file.data), keep_vba=keep_vba)
    try:
        ws = wb[upload.sheet.sheet_name]
        col = asset_id_column(upload)
        for r in result.rows:
            ws.cell(row=r.item.excel_row, column=col).value = _cell_value(r.asset_id)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()
