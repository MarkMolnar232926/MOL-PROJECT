"""Write the reconciled workbook (PLAN.md section 7.2).

Five sheets: SAP_Export_Completed, Physical_Inventory_Annotated, Discrepancies, Summary,
Decisions_Log. Values only (no formulas), Arial, frozen header, autofilter.
"""

from __future__ import annotations

import datetime as dt
import io
import math

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .loader import EXCEL_ROW, LoadedInput, SheetData
from .models import MatchStatus, ReconResult

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name=FONT)
ASSET_ID_FILL = "FFFFFFCC"  # pale yellow, used when the source has no fill of its own
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
MIN_WIDTH, MAX_WIDTH = 8, 60

STATUS_FILLS = {
    MatchStatus.MATCHED: "FFE2EFDA",
    MatchStatus.LOCATION_MISMATCH: "FFFFF2CC",
    MatchStatus.NEEDS_DECISION: "FFFCE4D6",
    MatchStatus.MANUAL: "FFDDEBF7",
    MatchStatus.PHYSICAL_ONLY: "FFF8CBAD",
    MatchStatus.SAP_ONLY: "FFF8CBAD",
    MatchStatus.DEFECTIVE: "FFEDEDED",
    MatchStatus.DUPLICATE: "FFEDEDED",
    MatchStatus.UNCLASSIFIED: "FFF8CBAD",
}

SAP_EXTRA = ["Match status", "Confidence", "Matched physical row", "QR agrees?", "Notes"]
PHYSICAL_EXTRA = ["Match status", "Matched SAP row", "SAP Item Name", "Notes"]


def export_filename(now: dt.datetime | None = None) -> str:
    return f"reconciled_{(now or dt.datetime.now()):%Y%m%d_%H%M%S}.xlsx"


def _cell_value(v):
    """Excel-safe value: no NaN/NaT, pandas timestamps as datetimes."""
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.to_pydatetime()
    return v


def _asset_id_value(asset_id: str | None):
    """Write numeric IDs as numbers (as in the physical sheet), anything else as text."""
    if asset_id is None:
        return None
    return int(asset_id) if asset_id.isdigit() else asset_id


def _utc_naive(t: dt.datetime) -> dt.datetime:
    if t.tzinfo is not None:
        t = t.astimezone(dt.UTC).replace(tzinfo=None)
    return t.replace(microsecond=0)


def _yes_no(v: bool | None) -> str | None:
    return None if v is None else ("Yes" if v else "No")


def _write_table(ws: Worksheet, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_cell_value(v) for v in row])
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            if isinstance(cell.value, dt.datetime):
                has_time = cell.value.time() != dt.time(0)
                cell.number_format = DATETIME_FORMAT if has_time else DATE_FORMAT
            elif isinstance(cell.value, dt.date):
                cell.number_format = DATE_FORMAT
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    _autosize(ws)


def _autosize(ws: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            text = f"{v:%Y-%m-%d}" if isinstance(v, dt.date) else str(v)
            widths[cell.column] = max(widths.get(cell.column, 0), len(text))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(
            MAX_WIDTH, max(MIN_WIDTH, width + 2)
        )


def _fill_column(ws: Worksheet, col: int, rgb: str, n_rows: int) -> None:
    fill = PatternFill("solid", fgColor=rgb)
    for r in range(2, n_rows + 2):
        ws.cell(row=r, column=col).fill = fill


def _fill_status(ws: Worksheet, col: int, statuses: list[MatchStatus]) -> None:
    for r, status in enumerate(statuses, start=2):
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=STATUS_FILLS[status])


def _original_rows(sheet: SheetData) -> list[tuple[int, list]]:
    cols = [c for c in sheet.frame.columns if c != EXCEL_ROW]
    return [(int(row[EXCEL_ROW]), [row[c] for c in cols]) for row in sheet.frame.to_dict("records")]


def _sap_sheet(ws: Worksheet, sheet: SheetData, result: ReconResult) -> None:
    by_row = {r.excel_row: r for r in result.sap_rows}
    asset_col = sheet.headers.index(sheet.column_map["asset_id"])
    rows, statuses = [], []
    for excel_row, values in _original_rows(sheet):
        r = by_row[excel_row]
        values = list(values)
        # Only final IDs are written; a pending tie slot stays empty (never the suggestion).
        values[asset_col] = _asset_id_value(r.asset_id)
        rows.append(
            values
            + [
                r.match_status.value,
                r.confidence.value if r.confidence else None,
                r.matched_physical_row,
                _yes_no(r.qr_agrees),
                " ".join(r.notes) or None,
            ]
        )
        statuses.append(r.match_status)
    _write_table(ws, sheet.headers + SAP_EXTRA, rows)
    fill = sheet.column_fills.get(sheet.headers[asset_col], ASSET_ID_FILL)
    _fill_column(ws, asset_col + 1, fill, len(rows))
    _fill_status(ws, len(sheet.headers) + 1, statuses)


def _physical_sheet(ws: Worksheet, sheet: SheetData, result: ReconResult) -> None:
    by_row = {r.excel_row: r for r in result.physical_rows}
    rows, statuses = [], []
    for excel_row, values in _original_rows(sheet):
        r = by_row[excel_row]
        rows.append(
            list(values)
            + [r.match_status.value, r.matched_sap_row, r.sap_type, " ".join(r.notes) or None]
        )
        statuses.append(r.match_status)
    _write_table(ws, sheet.headers + PHYSICAL_EXTRA, rows)
    _fill_status(ws, len(sheet.headers) + 1, statuses)


def _discrepancy_sheet(ws: Worksheet, result: ReconResult) -> None:
    phys = {r.excel_row: r for r in result.physical_rows}
    sap = {r.excel_row: r for r in result.sap_rows}
    rows = []
    for d in result.discrepancies:
        item = None
        if d.sap_row in sap:
            item = sap[d.sap_row].item_name
        elif d.physical_row in phys:
            item = phys[d.physical_row].sap_type or phys[d.physical_row].item_name
        rows.append(
            [d.kind.value, d.physical_row, d.sap_row, _asset_id_value(d.asset_id), item, d.message]
        )
    headers = [
        "Issue",
        "Physical_Inventory row",
        "SAP_Export row",
        "Asset ID",
        "Item",
        "Explanation",
    ]
    _write_table(ws, headers, rows)


def _summary_sheet(
    ws: Worksheet, loaded: LoadedInput, result: ReconResult, generated_at: dt.datetime
) -> None:
    s, qr = result.summary, result.qr_assessment

    def pct(v: float | None) -> str:
        return "n/a" if v is None else f"{v:.1f} %"

    rows: list[list] = [
        ["Run", "Generated at", generated_at.replace(microsecond=0)],
        [
            "Run",
            "Physical_Inventory source",
            f"{loaded.physical.file_name} › {loaded.physical.sheet_name}",
        ],
        ["Run", "SAP_Export source", f"{loaded.sap.file_name} › {loaded.sap.sheet_name}"],
        ["Run", "Rule-set version", s.rules_version],
        ["Run", "QR matching", "On" if s.use_qr else "Off"],
        ["Run", "Manual decisions", s.manual_decisions],
        ["Totals", "Physical rows", s.physical_total],
        ["Totals", "SAP rows", s.sap_total],
        ["Totals", "Pairs (incl. pending tie suggestions)", s.pairs],
        ["Totals", "Location mismatches", s.location_mismatches],
        ["Totals", "Tie groups", s.tie_groups],
        ["Totals", "Tie slots", s.tie_slots],
        ["Totals", "Tie slots still needing a decision", s.tie_slots_pending],
    ]
    rows += [["Physical status", k, v] for k, v in _ordered(s.physical_status_counts)]
    rows += [["SAP status", k, v] for k, v in _ordered(s.sap_status_counts)]
    rows += [["Confidence", k, v] for k, v in sorted(s.confidence_counts.items())]
    rows += [
        ["QR assessment", "QR column present", "Yes" if qr.column_present else "No"],
        [
            "QR assessment",
            "Parseable QR codes",
            f"{qr.parseable}/{qr.total_rows} ({pct(qr.parseable_pct)})",
        ],
        [
            "QR assessment",
            "QR ID exists in Physical_Inventory",
            f"{qr.present_in_physical}/{qr.total_rows} ({pct(qr.present_pct)})",
        ],
        [
            "QR assessment",
            "Agreement with attribute match (QR off)",
            f"{qr.agreeing}/{qr.compared} ({pct(qr.agreement_pct)})",
        ],
        [
            "QR assessment",
            "Agreement outside tie groups",
            f"{qr.agreeing_outside_ties}/{qr.compared_outside_ties} "
            f"({pct(qr.agreement_outside_ties_pct)})",
        ],
        ["QR assessment", "QR codes look reliable", "Yes" if qr.looks_reliable else "No"],
    ]
    rows += [["Warning", "", w] for w in result.warnings]
    _write_table(ws, ["Section", "Item", "Value"], rows)


def _ordered(counts: dict[str, int]) -> list[tuple[str, int]]:
    order = [m.value for m in MatchStatus]
    return sorted(counts.items(), key=lambda kv: order.index(kv[0]))


def _decisions_sheet(ws: Worksheet, result: ReconResult) -> None:
    sap = {r.excel_row: r for r in result.sap_rows}
    rows = [
        [
            d.group_id,
            d.sap_row,
            sap[d.sap_row].serial_no if d.sap_row in sap else None,
            _asset_id_value(d.asset_id) if d.asset_id else "none (left unmatched)",
            _asset_id_value(d.proposed_asset_id),
            _utc_naive(d.decided_at),
        ]
        for d in sorted(result.decisions, key=lambda d: (d.decided_at, d.sap_row))
    ]
    headers = [
        "Tie group",
        "SAP_Export row",
        "Serial No.",
        "Chosen Asset ID",
        "Proposed Asset ID",
        "Decided at (UTC)",
    ]
    _write_table(ws, headers, rows)


def export_workbook(
    loaded: LoadedInput, result: ReconResult, generated_at: dt.datetime | None = None
) -> bytes:
    """Build the reconciled workbook and return it as .xlsx bytes."""
    generated_at = generated_at or dt.datetime.now()
    wb = Workbook()
    wb.remove(wb.active)
    _sap_sheet(wb.create_sheet("SAP_Export_Completed"), loaded.sap, result)
    _physical_sheet(wb.create_sheet("Physical_Inventory_Annotated"), loaded.physical, result)
    _discrepancy_sheet(wb.create_sheet("Discrepancies"), result)
    _summary_sheet(wb.create_sheet("Summary"), loaded, result, generated_at)
    _decisions_sheet(wb.create_sheet("Decisions_Log"), result)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main(argv: list[str]) -> int:
    """python -m recon.export WORKBOOK.xlsx [-o OUT.xlsx] [--qr] [--accept-suggestions]"""
    import sys
    from pathlib import Path

    from .engine import accept_suggestions, reconcile
    from .errors import ReconError
    from .loader import InputFile, load_inputs
    from .normalize import normalize

    args = list(argv)
    use_qr = "--qr" in args
    accept = "--accept-suggestions" in args
    out = None
    if "-o" in args:
        i = args.index("-o")
        out = Path(args[i + 1])
        del args[i : i + 2]
    paths = [a for a in args if not a.startswith("--")]
    if len(paths) not in (1, 2):
        print(main.__doc__)
        return 2
    roles = [None] if len(paths) == 1 else ["physical", "sap"]
    try:
        data = normalize(
            load_inputs(
                [
                    InputFile(Path(p).name, Path(p).read_bytes(), r)
                    for p, r in zip(paths, roles, strict=True)
                ]
            )
        )
        result = reconcile(data, use_qr=use_qr)
        if accept:
            result = reconcile(data, use_qr=use_qr, decisions=accept_suggestions(result))
    except ReconError as exc:
        print(f"ERROR [{exc.code}]: {exc.message}", file=sys.stderr)
        return 1
    out = out or Path(export_filename())
    out.write_bytes(export_workbook(data.loaded, result))
    print(f"wrote {out} ({result.summary.tie_slots_pending} tie slots pending)")
    return 0


if __name__ == "__main__":
    import sys

    raise SystemExit(main(sys.argv[1:]))
